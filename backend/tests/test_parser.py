"""Tests for multi-format document parser and chunking logic."""

import io
from pathlib import Path
import docx
import fitz
import pytest
from fastapi import HTTPException
from backend.services.document_parser import DocumentParser

SAMPLE_DIR = Path(__file__).resolve().parent.parent / "sample_data"


def test_parse_valid_txt():
    """Verify parsing of valid legal plain text."""
    content = (SAMPLE_DIR / "commercial_lease_agreement.txt").read_bytes()
    doc_content = DocumentParser.parse_document(
        content=content,
        filename="commercial_lease_agreement.txt",
        ext=".txt",
        doc_id="test_doc_1",
        jurisdiction_country="United States",
        jurisdiction_state="New York",
    )

    assert doc_content.metadata.doc_id == "test_doc_1"
    assert doc_content.metadata.word_count > 500
    assert len(doc_content.chunks) > 0
    section_titles = [c.section_title for c in doc_content.chunks]
    assert any("LEASED PREMISES" in s or "BASE RENT" in s or "SECTION" in s for s in section_titles)


def test_parse_empty_txt_raises_400():
    """Verify empty text file throws 400 error."""
    with pytest.raises(HTTPException) as exc_info:
        DocumentParser.parse_txt(b"   \n\n   ", "empty.txt")
    assert exc_info.value.status_code == 400


def test_chunking_character_offsets():
    """Verify chunking maintains consistent char_start and char_end offsets."""
    content = b"SECTION 1. DEFINITIONS\nThis is clause A.\n\nSECTION 2. TERM\nThis is clause B."
    doc = DocumentParser.parse_document(
        content=content,
        filename="mini.txt",
        ext=".txt",
        doc_id="mini_1",
    )
    assert len(doc.chunks) >= 1
    for chunk in doc.chunks:
        assert chunk.char_end >= chunk.char_start
        assert chunk.page_number >= 1


def test_parse_in_memory_pdf():
    """Verify PyMuPDF parser extracts text and detects page breaks correctly."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "SECTION 1. COMMERCIAL LEASE AGREEMENT\nThis agreement is made between Landlord and Tenant.")
    pdf_bytes = doc.tobytes()
    doc.close()

    parsed = DocumentParser.parse_document(pdf_bytes, "sample.pdf", ".pdf", "pdf_1")
    assert parsed.metadata.page_count == 1
    assert "Landlord" in parsed.raw_text


def test_parse_in_memory_docx():
    """Verify python-docx parser extracts text paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("SECTION 1. CONSULTING TERMS", level=1)
    doc.add_paragraph("The Contractor shall render professional engineering advisory services.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Hourly Rate"
    table.rows[0].cells[1].text = "$150/hr"
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    parsed = DocumentParser.parse_document(docx_bytes, "sample.docx", ".docx", "docx_1")
    assert "Contractor" in parsed.raw_text
    assert "Hourly Rate" in parsed.raw_text
