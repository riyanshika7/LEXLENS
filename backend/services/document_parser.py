"""Multi-format document parsing engine supporting PDF, DOCX, and TXT with robust error recovery."""

import io
from typing import List, Tuple
from fastapi import HTTPException, status
import fitz  # PyMuPDF
import docx

from backend.schemas.document import DocumentMetadata, DocumentContent
from backend.services.chunker import chunk_document


class DocumentParser:
    """Parses legal documents into structured pages, sections, and indexed chunks."""

    @staticmethod
    def parse_txt(content: bytes, filename: str) -> Tuple[str, List[Tuple[int, str]]]:
        """Decode TXT bytes into full text and page tuples."""
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = content.decode("latin-1")
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Unable to decode text file. Ensure UTF-8 encoding. Error: {str(e)}",
                )

        if not text.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded text document is completely empty or contains only whitespace.",
            )

        raw_pages = text.split("\x0c") if "\x0c" in text else []
        if not raw_pages:
            lines = text.splitlines(keepends=True)
            pages: List[Tuple[int, str]] = []
            cur_page_lines = []
            cur_char_count = 0
            page_num = 1

            for line in lines:
                cur_page_lines.append(line)
                cur_char_count += len(line)
                if cur_char_count >= 3000:
                    pages.append((page_num, "".join(cur_page_lines)))
                    cur_page_lines = []
                    cur_char_count = 0
                    page_num += 1

            if cur_page_lines:
                pages.append((page_num, "".join(cur_page_lines)))
        else:
            pages = [(i + 1, p) for i, p in enumerate(raw_pages) if p.strip()]

        if not pages:
            pages = [(1, text)]

        return text, pages

    @staticmethod
    def parse_pdf(content: bytes, filename: str) -> Tuple[str, List[Tuple[int, str]]]:
        """Extract text from PDF using PyMuPDF with corruption & password detection."""
        try:
            doc = fitz.open(stream=content, filetype="pdf")
        except Exception:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="We couldn't parse this PDF because the file structure is corrupted or unreadable. Please check the file and try again.",
            )

        if doc.is_encrypted:
            doc.close()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="We couldn't analyze this document because the file appears to be password protected or encrypted. Try uploading an unlocked copy.",
            )

        if doc.page_count == 0:
            doc.close()
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The PDF contains zero pages. Please upload a valid document.",
            )

        pages: List[Tuple[int, str]] = []
        full_text_parts: List[str] = []

        for page_idx in range(doc.page_count):
            page = doc.load_page(page_idx)
            page_text = page.get_text("text") or ""
            pages.append((page_idx + 1, page_text))
            full_text_parts.append(page_text)

        doc.close()
        full_text = "\n\n".join(full_text_parts).strip()

        if not full_text or len(full_text.strip()) < 20:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="No readable text found in this PDF. This appears to be a scanned image-only PDF without an embedded OCR text layer. Please convert or run OCR on the document before uploading.",
            )

        return full_text, pages

    @staticmethod
    def parse_docx(content: bytes, filename: str) -> Tuple[str, List[Tuple[int, str]]]:
        """Extract text and tables from DOCX document."""
        try:
            doc_file = io.BytesIO(content)
            doc = docx.Document(doc_file)
        except Exception:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Unable to read DOCX file. The archive may be corrupted or created with an unsupported Word format.",
            )

        full_text_parts: List[str] = []

        for p in doc.paragraphs:
            if p.text.strip():
                full_text_parts.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    full_text_parts.append(" | ".join(row_text))

        full_text = "\n\n".join(full_text_parts).strip()
        if not full_text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="The uploaded DOCX document contains no text content.",
            )

        pages: List[Tuple[int, str]] = []
        cur_chars = 0
        cur_paragraphs = []
        page_num = 1

        for part in full_text_parts:
            cur_paragraphs.append(part)
            cur_chars += len(part)
            if cur_chars >= 3000:
                pages.append((page_num, "\n\n".join(cur_paragraphs)))
                cur_paragraphs = []
                cur_chars = 0
                page_num += 1

        if cur_paragraphs:
            pages.append((page_num, "\n\n".join(cur_paragraphs)))

        return full_text, pages

    @classmethod
    def parse_document(
        cls,
        content: bytes,
        filename: str,
        ext: str,
        doc_id: str,
        jurisdiction_country: str = "Unspecified",
        jurisdiction_state: str = "Unspecified",
    ) -> DocumentContent:
        """Parse raw document bytes into complete DocumentContent object."""
        ext_clean = ext.lower().strip()
        if ext_clean == ".pdf":
            raw_text, pages = cls.parse_pdf(content, filename)
        elif ext_clean == ".docx":
            raw_text, pages = cls.parse_docx(content, filename)
        elif ext_clean == ".txt":
            raw_text, pages = cls.parse_txt(content, filename)
        else:
            raise HTTPException(
                status_code=status.HTTP_415_UNSUPPORTED_MEDIA_TYPE,
                detail=f"Unsupported format: {ext_clean}. Supported formats: .pdf, .docx, .txt",
            )

        chunks = chunk_document(pages)
        word_count = len(raw_text.split())

        metadata = DocumentMetadata(
            doc_id=doc_id,
            filename=filename,
            file_type=ext_clean.replace(".", ""),
            size_bytes=len(content),
            page_count=len(pages),
            word_count=word_count,
            jurisdiction_country=jurisdiction_country or "Unspecified",
            jurisdiction_state=jurisdiction_state or "Unspecified",
        )

        return DocumentContent(
            metadata=metadata,
            raw_text=raw_text,
            chunks=chunks,
        )
